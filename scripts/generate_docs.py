#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FOLIO — Generador de la documentación extendida en Word.

Lee DOCUMENTACION_ARQUITECTURA.md (fuente de verdad), lo expande con contexto,
diagramas en texto, decisiones de arquitectura, guía para developers, glosario y
apéndices, y genera DOCUMENTACION_ARQUITECTURA_COMPLETA.docx en la raíz del repo.

Uso:
    pip install python-docx
    python scripts/generate_docs.py
    python scripts/generate_docs.py --pending "mensaje del commit"   # añade el commit en curso al apéndice E

Al terminar verifica la coherencia .md <-> .docx (encabezados, migraciones SQL, términos de auth)
y sale con código 1 si hay divergencias.

El .md sigue siendo la fuente de verdad: NO edites el .docx a mano, regéneralo.
Los apéndices A (tablas creadas por SQL), B (firmas de RPC) y D (fechas de migraciones)
se contrastan con supabase/*.sql y con `git log` cuando están disponibles.
"""
import datetime
import math
import re
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("Falta python-docx. Instálalo con:  pip install python-docx")

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "DOCUMENTACION_ARQUITECTURA.md"
OUT_PATH = ROOT / "DOCUMENTACION_ARQUITECTURA_COMPLETA.docx"
SQL_DIR = ROOT / "supabase"

DOC_VERSION = "1.2 (extendida)"
PENDING_COMMIT = ""   # se rellena con --pending "mensaje" para incluir el commit en curso en el apéndice E
AUTHOR = "FOLIO Team"
SERIF = "Georgia"
SANS = "Calibri"
MONO = "Consolas"
ACCENT = RGBColor(0x7A, 0x2E, 0x2E)   # accent de FOLIO (theme.js)
GRAY = RGBColor(0x55, 0x55, 0x55)
CODE_FILL = "F0F0F0"
NOTE_FILL = "FBF3E4"
HEAD_FILL = "E9DFCB"
TEXT_WIDTH_IN = 6.5

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def fecha_es(d):
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


# ============================================================================
# 1. Helpers XML / estilos
# ============================================================================
def set_font(run_or_style, name, size=None, bold=None, italic=None, color=None):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style.element, "get_or_add_rPr") else run_or_style.element.rPr
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def para_borders(ppr, sides=("top", "left", "bottom", "right"), color="BBBBBB", sz=4, left_sz=None, left_color=None):
    pbdr = OxmlElement("w:pBdr")
    for side in sides:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        s = left_sz if (side == "left" and left_sz) else sz
        b.set(qn("w:sz"), str(s))
        b.set(qn("w:space"), "4")
        b.set(qn("w:color"), left_color if (side == "left" and left_color) else color)
        pbdr.append(b)
    ppr.append(pbdr)


def add_field(paragraph, instr, placeholder=""):
    """Campo complejo de Word (PAGE, TOC, ...)."""
    def fc(t):
        r = paragraph.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)
        return r
    fc("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    if placeholder:
        paragraph.add_run(placeholder)
    fc("end")


def setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    set_font(normal, SANS, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    sizes = {1: 20, 2: 15, 3: 12.5}
    for lvl, size in sizes.items():
        h = st[f"Heading {lvl}"]
        set_font(h, SERIF, size, bold=True, italic=False, color=ACCENT if lvl == 1 else RGBColor(0x33, 0x33, 0x33))
        h.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
        h.paragraph_format.space_after = Pt(8 if lvl == 1 else 4)
        h.paragraph_format.keep_with_next = True
        if lvl == 1:
            h.paragraph_format.page_break_before = True
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        set_font(st[name], SANS, 10.5)
        st[name].paragraph_format.space_after = Pt(2)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.different_first_page_header_footer = True   # portada sin número

    # Pie con numeración de página
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("FOLIO · Documentación de arquitectura  —  Página ")
    set_font(r, SANS, 8.5, color=GRAY)
    add_field(fp, "PAGE", "1")
    for run in fp.runs:
        set_font(run, SANS, 8.5, color=GRAY)

    # Que Word ofrezca actualizar la tabla de contenidos al abrir
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)

    cp = doc.core_properties
    cp.title = "FOLIO — Documentación de arquitectura técnica completa"
    cp.author = AUTHOR
    cp.subject = "Arquitectura, base de datos, servicios y guía de desarrollo"
    cp.language = "es-MX"


# ============================================================================
# 2. Builder
# ============================================================================
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?![\s*])[^*\n]+?(?<!\s)\*(?![\w*]))")


class Builder:
    def __init__(self, doc, toc_entries=None):
        self.doc = doc
        self.toc_entries = toc_entries or []
        self.headings = []          # (level, texto numerado) para el TOC estático
        self.c = [0, 0, 0]
        self.app = None             # letra de apéndice activo

    # ---- texto en línea ----
    def inline(self, par, text, size=None, bold_all=False, color=None):
        for part in INLINE_RE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                for sub in re.split(r"(`[^`]+`)", part[2:-2]):   # `código` dentro de negrita
                    if not sub:
                        continue
                    if sub.startswith("`") and sub.endswith("`") and len(sub) > 2:
                        r = par.add_run(sub[1:-1])
                        set_font(r, MONO, (size or 10.5) - 1.5, bold=True, color=RGBColor(0x8A, 0x2B, 0x2B))
                        shade(r._r.get_or_add_rPr(), "F0F0F0")
                    else:
                        r = par.add_run(sub)
                        set_font(r, SANS, size, bold=True, color=color)
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                r = par.add_run(part[1:-1])
                set_font(r, MONO, (size or 10.5) - 1.5, color=RGBColor(0x8A, 0x2B, 0x2B))
                shade(r._r.get_or_add_rPr(), "F0F0F0")
            elif len(part) > 2 and part.startswith("*") and part.endswith("*") and not part.startswith("**"):
                r = par.add_run(part[1:-1])
                set_font(r, SANS, size, bold=bold_all or None, italic=True, color=color)
            else:
                r = par.add_run(part)
                set_font(r, SANS, size, bold=bold_all or None, color=color)

    # ---- encabezados ----
    def heading(self, level, text, prefix=None):
        if prefix is None:
            if self.app:
                if level == 1:
                    prefix = ""
                elif level == 2:
                    self.c[1] += 1
                    self.c[2] = 0
                    prefix = f"{self.app}.{self.c[1]}  "
                else:
                    self.c[2] += 1
                    prefix = f"{self.app}.{self.c[1]}.{self.c[2]}  "
            else:
                if level == 1:
                    self.c = [self.c[0] + 1, 0, 0]
                    prefix = f"{self.c[0]}.  "
                elif level == 2:
                    self.c[1] += 1
                    self.c[2] = 0
                    prefix = f"{self.c[0]}.{self.c[1]}  "
                else:
                    self.c[2] += 1
                    prefix = f"{self.c[0]}.{self.c[1]}.{self.c[2]}  "
        full = f"{prefix}{text}".replace("`", "")
        self.doc.add_heading(full, level=level)
        self.headings.append((level, full))

    def begin_appendix(self, letter, title):
        self.app = letter
        self.c = [0, 0, 0]
        self.heading(1, f"Apéndice {letter}. {title}", prefix="")

    # ---- bloques ----
    def paragraph(self, text):
        p = self.doc.add_paragraph()
        self.inline(p, text)
        return p

    def bullets(self, items, level=0):
        style = ["List Bullet", "List Bullet 2", "List Bullet 3"][min(level, 2)]
        for it in items:
            p = self.doc.add_paragraph(style=style)
            self.inline(p, it)

    def ordered(self, items, start=1):
        for i, it in enumerate(items, start):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(3)
            self.inline(p, f"{i}. {it}")

    def code(self, text):
        p = self.doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        para_borders(ppr, color="CCCCCC")
        shade(ppr, CODE_FILL)
        pf = p.paragraph_format   # después de pBdr/shd para respetar el orden del esquema
        pf.left_indent = Inches(0.1)
        pf.right_indent = Inches(0.1)
        pf.space_before = Pt(4)
        pf.space_after = Pt(8)
        pf.line_spacing = 1.0
        lines = text.rstrip("\n").split("\n")
        for i, ln in enumerate(lines):
            r = p.add_run(ln.replace("\t", "    ") or " ")
            set_font(r, MONO, 8)
            if i < len(lines) - 1:
                r.add_break(WD_BREAK.LINE)

    def note(self, text):
        p = self.doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        para_borders(ppr, sides=("left",), left_sz=24, left_color="7A2E2E")
        shade(ppr, NOTE_FILL)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.15)
        pf.space_before = Pt(4)
        pf.space_after = Pt(8)
        self.inline(p, text, size=10)

    def table(self, header, rows, widths=None):
        ncols = len(header)
        if widths is None:
            weights = []
            for ci in range(ncols):
                m = max([len(str(header[ci]))] + [len(str(r[ci])) if ci < len(r) else 0 for r in rows])
                weights.append(max(6.0, min(math.sqrt(m) * 3.2, 26.0)))
            tot = sum(weights)
            widths = [w / tot for w in weights]
        tbl = self.doc.add_table(rows=1, cols=ncols)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        # cabecera repetida
        trpr = tbl.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for ci, h in enumerate(header):
            cell = tbl.rows[0].cells[ci]
            cell.width = Inches(TEXT_WIDTH_IN * widths[ci])
            shade(cell._tc.get_or_add_tcPr(), HEAD_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            self.inline(p, str(h), size=8.5, bold_all=True)
        for row in rows:
            cells = tbl.add_row().cells
            for ci in range(ncols):
                txt = str(row[ci]) if ci < len(row) else ""
                cells[ci].width = Inches(TEXT_WIDTH_IN * widths[ci])
                p = cells[ci].paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                self.inline(p, txt, size=8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def render(self, blocks):
        for b in blocks:
            k = b[0]
            if k == "h2":
                self.heading(2, b[1])
            elif k == "h3":
                self.heading(3, b[1])
            elif k == "p":
                self.paragraph(b[1])
            elif k == "ul":
                self.bullets(b[1])
            elif k == "ol":
                self.ordered(b[1])
            elif k == "code":
                self.code(b[1])
            elif k == "note":
                self.note(b[1])
            elif k == "table":
                self.table(b[1], b[2], b[3] if len(b) > 3 else None)
            else:
                raise ValueError(f"bloque desconocido: {k}")

    # ---- portada + TOC ----
    def cover(self):
        d = self.doc
        for _ in range(5):
            d.add_paragraph()
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("FOLIO")
        set_font(r, SERIF, 54, bold=True, color=ACCENT)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Documentación de arquitectura técnica completa")
        set_font(r, SERIF, 20, color=RGBColor(0x33, 0x33, 0x33))
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Biblioteca personal y red social de lectura gamificada (PWA · React + Supabase)")
        set_font(r, SANS, 12, italic=True, color=GRAY)
        for _ in range(6):
            d.add_paragraph()
        for label, value in (("Versión", DOC_VERSION), ("Fecha", fecha_es(datetime.date.today())), ("Autor", AUTHOR)):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"{label}:  ")
            set_font(r, SANS, 11, bold=True, color=GRAY)
            r = p.add_run(value)
            set_font(r, SANS, 11)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(40)
        r = p.add_run("Documento generado desde DOCUMENTACION_ARQUITECTURA.md (fuente de verdad) con "
                      "scripts/generate_docs.py. No editar a mano: regenerar.")
        set_font(r, SANS, 9, italic=True, color=GRAY)

    def toc(self):
        d = self.doc
        h = d.add_paragraph()
        h.paragraph_format.page_break_before = True
        r = h.add_run("Tabla de contenidos")
        set_font(r, SERIF, 20, bold=True, color=ACCENT)
        n = d.add_paragraph()
        r = n.add_run("Si los números de página no aparecen, haz clic derecho sobre la tabla → «Actualizar campo» "
                      "(Word también lo ofrece al abrir el documento).")
        set_font(r, SANS, 8.5, italic=True, color=GRAY)
        entries = [(lv, t) for lv, t in self.toc_entries if lv <= 2] or [(1, "(actualiza el campo para generar el índice)")]
        for i, (lv, text) in enumerate(entries):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.0 if lv == 1 else 0.3)
            p.paragraph_format.space_after = Pt(1 if lv == 2 else 3)
            if i == 0:
                self._fld(p, "begin")
                self._instr(p, 'TOC \\o "1-3" \\h \\z \\u')
                self._fld(p, "separate")
            r = p.add_run(text)
            set_font(r, SANS, 10 if lv == 1 else 9, bold=(lv == 1) or None)
            if i == len(entries) - 1:
                self._fld(p, "end")

    @staticmethod
    def _fld(p, t):
        r = p.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)

    @staticmethod
    def _instr(p, text):
        r = p.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {text} "
        r._r.append(it)


# ============================================================================
# 3. Parser de Markdown (subset usado por DOCUMENTACION_ARQUITECTURA.md)
# ============================================================================
def split_sections(md_text):
    """Devuelve (preámbulo, [ {title, roman, lines} ]) partiendo por '## '."""
    pre, sections, cur = [], [], None
    for line in md_text.splitlines():
        if line.startswith("## "):
            cur = {"title": line[3:].strip(), "lines": []}
            sections.append(cur)
        elif cur is None:
            pre.append(line)
        else:
            cur["lines"].append(line)
    for s in sections:
        m = re.match(r"^([IVX]+)\.\s+(.*)$", s["title"])
        s["roman"], s["name"] = (m.group(1), m.group(2)) if m else (None, s["title"])
    return pre, sections


def nice_title(name):
    """'VISIÓN GENERAL' → 'Visión general' (conserva paréntesis tal cual)."""
    main, sep, rest = name.partition(" (")
    if main.isupper():
        main = main[0] + main[1:].lower()
    return main + (sep + rest if sep else "")


def parse_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    ncols = len(rows[0])
    for r in rows:
        if len(r) > ncols:   # pipe suelto dentro de una celda: se fusiona con la última
            r[ncols - 1:] = [" | ".join(r[ncols - 1:])]
    header, body = rows[0], [r for r in rows[1:] if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    return header, body


def parse_blocks(lines):
    blocks, i, n = [], 0, len(lines)
    list_re = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")
    while i < n:
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
        elif re.match(r"^#{3,4}\s", line):
            lvl = 2 if line.startswith("### ") else 3
            blocks.append(("h2" if lvl == 2 else "h3", re.sub(r"^\d+\.\s+", "", line.lstrip("#").strip())))
            i += 1
        elif s.startswith("|"):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append(lines[i])
                i += 1
            header, body = parse_table(buf)
            blocks.append(("table", header, body))
        elif s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("note", " ".join(buf)))
        elif list_re.match(line):
            items = []   # (indent_level, ordered, text)
            while i < n:
                m = list_re.match(lines[i])
                if m:
                    indent = len(m.group(1).replace("\t", "    "))
                    items.append([min(indent // 3, 2), m.group(2)[0].isdigit(), m.group(3), m.group(2).rstrip(".")])
                    i += 1
                elif lines[i].strip() and lines[i].startswith("  ") and items:   # continuación
                    items[-1][2] += " " + lines[i].strip()
                    i += 1
                else:
                    break
            for lvl, ordered, text, num in items:
                blocks.append(("li", lvl, ordered, text, num))
        elif not s or s == "---":
            i += 1
        else:
            buf = []
            while i < n and lines[i].strip() and not lines[i].strip().startswith(("```", "|", ">", "#", "---")) \
                    and not list_re.match(lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("p", " ".join(buf)))
    return blocks


def render_md_blocks(b, blocks):
    """Como Builder.render pero convirtiendo los ('li', ...) en listas."""
    for blk in blocks:
        if blk[0] == "li":
            _, lvl, ordered, text, num = blk
            if ordered and lvl == 0:
                b.ordered([text], start=int(num))
            else:
                b.bullets([text], level=lvl)
        else:
            b.render([blk])


# ============================================================================
# 4. Datos desde el repo (SQL, git)
# ============================================================================
def strip_sql_comments(sql):
    return "\n".join(re.sub(r"--.*$", "", ln) for ln in sql.splitlines())


def parse_sql_tables():
    """{tabla: [(columna, definición)]} para los CREATE TABLE de supabase/*.sql."""
    tables = {}
    for f in sorted(SQL_DIR.glob("*.sql")):
        sql = strip_sql_comments(f.read_text(encoding="utf-8", errors="ignore"))
        for m in re.finditer(r"CREATE TABLE(?: IF NOT EXISTS)?\s+(?:public\.)?(\w+)\s*\(", sql):
            name, depth, j = m.group(1), 1, m.end()
            start = j
            while j < len(sql) and depth:
                depth += {"(": 1, ")": -1}.get(sql[j], 0)
                j += 1
            body = sql[start:j - 1]
            parts, buf, d = [], "", 0
            for ch in body:
                if ch == "(":
                    d += 1
                elif ch == ")":
                    d -= 1
                if ch == "," and d == 0:
                    parts.append(buf.strip())
                    buf = ""
                else:
                    buf += ch
            if buf.strip():
                parts.append(buf.strip())
            cols = []
            for p in parts:
                p = " ".join(p.split())
                if not p:
                    continue
                head = p.split(" ", 1)
                if head[0].upper() in ("PRIMARY", "UNIQUE", "CONSTRAINT", "CHECK", "FOREIGN"):
                    cols.append(("(restricción)", p))
                else:
                    cols.append((head[0], head[1] if len(head) > 1 else ""))
            tables[name] = (f.name, cols)
    return tables


def parse_sql_functions():
    funcs = set()
    for f in SQL_DIR.glob("*.sql"):
        for m in re.finditer(r"CREATE (?:OR REPLACE )?FUNCTION\s+(?:public\.)?(\w+)\s*\(",
                             f.read_text(encoding="utf-8", errors="ignore")):
            funcs.add(m.group(1))
    return funcs


def git(*args):
    try:
        out = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=20)
        return out.stdout.strip() if out.returncode == 0 else ""
    except Exception:
        return ""


def git_first_date(path):
    out = git("log", "--diff-filter=A", "--format=%ad", "--date=short", "--", str(path.relative_to(ROOT)))
    return out.splitlines()[-1] if out else ""


# ============================================================================
# 5. CONTENIDO EXTENDIDO
# ============================================================================
INTRO = [
    ("h2", "Qué es FOLIO"),
    ("p", "**FOLIO** es una aplicación web progresiva (PWA) en español (es-MX) que combina tres cosas: una **biblioteca personal** "
          "(qué libros tengo, cuáles leo, cuáles quiero, qué opino de ellos), una **red social de lectura** (feed, amigos, chat 1:1, "
          "citas compartidas, compatibilidad lectora) y una capa de **gamificación** inspirada en Duolingo: rachas diarias, una "
          "mascota virtual que sube de nivel con XP, gemas, logros y un resumen mensual tipo «Wrapped»."),
    ("p", "La propuesta de valor es convertir el hábito de leer —solitario y difícil de sostener— en algo visible, social y recompensado, "
          "sin castigar: la racha se **pausa**, nunca se resetea. Sobre esa base FOLIO añade recomendaciones (tres motores distintos, "
          "incluido un recomendador colaborativo entrenado offline) y, desde septiembre de 2026, el **Trueque**: intercambio de libros "
          "físicos entre usuarios de la misma zona de la Ciudad de México."),
    ("p", "Técnicamente es una SPA de React 18 servida por Vite, con **Supabase** como backend completo (Auth, Postgres con RLS y "
          "funciones SECURITY DEFINER como «servidor de autoridad») y un puñado de **funciones serverless de Vercel** que solo hacen de "
          "proxy (Anthropic, Google Books) y de lectura de recomendaciones. No existe un backend propio persistente."),
    ("h2", "Historia del proyecto"),
    ("p", "La historia se reconstruye del historial de git (los commits son la única bitácora formal del proyecto):"),
    ("table", ["Etapa", "Periodo (git)", "Qué ocurrió"], [
        ["MVP", "2026-05-04 → 2026-06-01",
         "Commit inicial y construcción de la app: feed social (con varias rondas de fixes de carga y race conditions), sistema de "
         "animación y accesibilidad, gemas y avatares, lector interno de cuentos, tutorial de primer login y rate limiting del proxy de Anthropic."],
        ["Auth real y RLS", "2026-07-02 →",
         "Migración a Supabase Auth con RLS en todas las tablas (`auth_rls_migration.sql`), tabla de recomendaciones y `features_migration.sql` "
         "(listas de usuario, borrado de cuenta, perfil público/privado, fecha de lectura flexible)."],
        ["Server authority", "2026-07-16",
         "XP y gemas pasan a decidirse en Postgres (`sprint1_server_authority.sql`: RPCs, triggers, ledger idempotente). Se extrae la "
         "capa de servicios y algunos componentes del monolito App.jsx."],
        ["Sprint de seguridad", "2026-09-13 → 2026-09-17",
         "Fixes de racha y caché de perfil; parches contra farmeo de recompensas, logros y rachas falsificables; privacidad real "
         "(`privacy_hardening.sql`); zona horaria unificada; proxy Anthropic con JWT. El commit se nombró «pre-INDAUTOR»."],
        ["Audit + Trueque", "2026-09-18 → 2026-09-19",
         "Sincronización de la documentación con el código, fixes de portada/ErrorBoundary/.env.example y el MVP del Trueque de libros "
         "de punta a punta (DB, RPCs, UI, docs)."],
        ["Fix de signup", "2026-09-18 (changelog)",
         "Se corrige el bug «permission denied for table users» del registro (INSERT plano, policies y GRANT por columna) y se añade "
         "esta versión extendida de la documentación en Word."],
        ["Fix real de cuentas huérfanas", "2026-09-19 (changelog)",
         "El fix anterior solo cubría de forma efectiva el signup. Ahora `ensureUserProfile` repara el perfil faltante en cada login, "
         "sesión guardada y evento `onAuthStateChange`, sin fallar en silencio, y `repair_all_orphan_profiles.sql` repara el histórico."],
        ["Diagnóstico en producción", "2026-09-19 (changelog)",
         "`ensureUserProfile` no dejaba rastro en producción y el perfil seguía sin crearse. Se añade instrumentación `[auth-debug]`, "
         "llamadas explícitas en login/registro/listener, un fallback defensivo al montar `App` y el log de versión del build."],
    ], [0.17, 0.2, 0.63]),
    ("h2", "Estado actual"),
    ("ul", [
        "**Producción:** proyecto Vercel `folio-final`. El repositorio GitHub **no** está conectado a Vercel: los deploys son manuales (`vercel --prod`); un `git push` no despliega.",
        "**Backend:** Supabase con RLS en todas las tablas. Las migraciones SQL se corren a mano en el SQL Editor (no hay CLI configurada).",
        "**Seguridad:** XP, nivel, gemas, logros y racha solo se escriben desde Postgres; `is_public` se hace cumplir en RLS; `users.email` está fuera del SELECT de clientes.",
        "**Trueque:** MVP funcional. No hay pago real de Folio Plus (`is_premium` se activa a mano).",
        "**Bug abierto (2026-09-19):** el auto-perfil de cuentas huérfanas no dejó rastro en producción; hay instrumentación `[auth-debug]` activa (`AUTH_DEBUG` en `authService.js`) a la espera de las trazas de un registro nuevo y de un login.",
        "**Calidad:** sin suite de tests; App.jsx sigue siendo un monolito de ~14.900 líneas (ver sección de deuda técnica).",
    ]),
    ("h2", "Roadmap"),
    ("p", "El roadmap sale de la sección de deuda técnica y de los TODOs de Sprint 2 del documento fuente; no es un compromiso de fechas."),
    ("table", ["Prioridad", "Item", "Origen"], [
        ["Corto plazo", "Correr `fix_signup_permissions.sql` y `repair_all_orphan_profiles.sql` en Supabase y desplegar (`vercel --prod`)", "Fix de huérfanas"],
        ["Corto plazo", "Conectar GitHub ↔ Vercel para deploy automático", "Deuda #3"],
        ["Corto plazo", "Mover `buyExtraSaves` y `daily_save_limits` a RPCs SECURITY DEFINER", "Deuda #4 y #5"],
        ["Medio plazo", "Suite mínima de tests (servicios y RPCs críticas) y code-splitting del bundle (1.2 MB)", "Deuda #6 y #8"],
        ["Medio plazo", "Seguir extrayendo vistas de App.jsx a `src/components/`", "Deuda #1"],
        ["Medio plazo", "Pago real de Folio Plus; push notifications de matches y mensajes", "Trueque Sprint 2"],
        ["Largo plazo", "Trueque: geolocalización real (Google Places), «Embajadas FOLIO», moderación (reportar/bloquear)", "Trueque Sprint 2"],
        ["Largo plazo", "Más mascotas (sistema evolutivo) y más motores de recomendación", "Puntos de extensión"],
    ], [0.14, 0.62, 0.24]),
    ("h2", "Cómo leer este documento"),
    ("ul", [
        "Los capítulos 2 a 14 siguen el documento fuente (`DOCUMENTACION_ARQUITECTURA.md`) sección por sección; cada uno abre con un apartado «Contexto adicional» con ejemplos y diagramas en texto.",
        "El capítulo 15 explica **por qué** se tomaron las decisiones de arquitectura, el 16 es la guía práctica para developers nuevos y el 17 un glosario.",
        "Los apéndices reúnen referencias: esquema de tablas (A), RPCs (B), variables de entorno (C) e historial de migraciones (D).",
        "Convención de marcas: ⚠️ = riesgo o deuda conocida; ✅ = resuelto.",
    ]),
]

DECISIONS = [
    ("note", "**Nota de método.** FOLIO no tiene un registro formal de decisiones (ADR). Las razones de este capítulo se reconstruyen "
             "del código, de la documentación y de las restricciones del producto; donde una decisión fue más histórica que deliberada "
             "(por ejemplo el monolito App.jsx) se dice explícitamente."),
    ("h2", "Supabase en lugar de Firebase"),
    ("ul", [
        "**Decisión:** Supabase (Postgres + Auth + Storage + RLS + RPCs) como backend completo.",
        "**Por qué:** el dominio es fuertemente relacional (usuarios, libros, amistades, posts, comentarios, matches del Trueque) y varias reglas críticas son consultas con joins y transacciones (matching, ledger de recompensas). Postgres las expresa de forma nativa; en Firestore habría que desnormalizar y duplicar datos. RLS permite que el cliente hable directo con la base con seguridad declarativa, y SECURITY DEFINER da un «servidor de autoridad» sin operar un servidor.",
        "**Alternativas:** Firebase (Firestore + Functions), backend propio (Express + Postgres).",
        "**Consecuencias:** las reglas de negocio viven en SQL (más difícil de testear y versionar; migraciones manuales), y hay que pensar en GRANTs y policies como parte del modelo de seguridad — el bug de signup de septiembre 2026 es un ejemplo de ese costo.",
    ]),
    ("h2", "React + Vite (SPA) en lugar de Next.js"),
    ("ul", [
        "**Decisión:** SPA de React 18 con Vite 5, sin SSR ni router de framework.",
        "**Por qué:** la app es casi por completo post-login, por lo que el SEO es irrelevante; el público es mobile-first y la app se instala como PWA. Una SPA estática es más simple de desplegar, cachear y servir offline con Workbox. No hay páginas públicas que renderizar en servidor.",
        "**Alternativas:** Next.js (App Router), Remix.",
        "**Consecuencias:** no hay routing por URL (la navegación es un estado `tab`; los deep-links se leen de `window.location.search`), lo que complica compartir enlaces a vistas internas. Si algún día se necesitan páginas públicas indexables (perfiles, listas), habrá que añadirlas aparte o migrar.",
    ]),
    ("h2", "Vercel como hosting y proxy"),
    ("ul", [
        "**Decisión:** Vercel sirve el bundle estático y las funciones `api/*` (proxies).",
        "**Por qué:** despliegue de un SPA + funciones serverless con una sola herramienta; las funciones ocultan la API key de Anthropic y evitan CORS con Google Books. Como solo hay proxies sin estado, no se necesita un servidor persistente.",
        "**Alternativas:** Netlify, Cloudflare Pages/Workers, Supabase Edge Functions (Deno).",
        "**Consecuencias:** el rate limit del proxy es un `Map` en memoria que se reinicia con cada cold start. Además, hoy los deploys son manuales porque la integración Git↔Vercel no está conectada.",
    ]),
    ("h2", "El monolito App.jsx (~14.900 líneas)"),
    ("ul", [
        "**Decisión:** casi todas las vistas y helpers viven en un solo archivo.",
        "**Por qué:** honestamente, es una decisión **histórica más que deliberada**: el proyecto empezó como MVP de un solo archivo y creció por iteración rápida. Tuvo una ventaja real: cualquier cambio se busca con grep en un solo lugar, sin saltar entre módulos.",
        "**Consecuencias:** riesgo de conflictos de merge, builds y HMR más lentos, dificultad para testear y bugs al mover código (el caso `fetchStreakData` no exportado). La auditoría de fase 1 puntuó la arquitectura de código con 3/10 por este motivo.",
        "**Camino de salida:** ya iniciado con la capa `src/services/` y `src/components/`; el patrón es mover una vista con imports de servicios y props explícitas (commit `3b3adfd`) verificando exports con grep.",
    ]),
    ("h2", "Sin router y sin state manager"),
    ("p", "La navegación es un estado `tab` y el estado global se reparte con `useState`/`useEffect` más prop drilling desde `MainApp`, complementado con **event buses** caseros (`petBus`, `gemsEventBus`, `achievementBus`) para que los servicios notifiquen a la UI sin acceso a React. Es suficiente para una app de una sola pantalla principal con tabs y evita dependencias; el coste es prop drilling profundo y un `MainApp` con ~40 estados."),
    ("h2", "Estilos inline con `palette` mutable"),
    ("p", "`theme.js` exporta un objeto mutable `palette` que `MainApp` reasigna en cada render según el tema (claro/oscuro/sistema). Los componentes leen `palette.xxx` en estilos inline. Ventaja: modo oscuro instantáneo sin CSS variables ni contexto. Desventaja: los estilos no son cacheables ni sobrescribibles por CSS y Tailwind se usa solo de forma parcial."),
    ("h2", "Autoridad en el servidor (Postgres) para todo lo «ganable»"),
    ("p", "XP, nivel, gemas, logros y racha nunca se escriben desde el cliente: solo triggers y RPCs SECURITY DEFINER que llaman a `folio_award`, con un ledger idempotente (`reward_ledger`). Se eligió Postgres y no funciones serverless porque las reglas necesitan atomicidad con los datos (un trigger sobre `books` se ejecuta en la misma transacción que el UPDATE) y porque así el cliente no puede saltarse la validación. Es la base anti-cheat de la economía."),
    ("h2", "PWA en lugar de app nativa"),
    ("p", "Una PWA instalable con Workbox (autoUpdate + skipWaiting) llega a Android e iOS sin tiendas de apps, se actualiza al instante y comparte el 100% del código. Se paga con menos acceso a APIs nativas (las notificaciones programadas se resuelven con `sw-notifications.js`) y con el cuidado de `vercel.json`: `Cache-Control: no-store` en `index.html` y `sw.js` es imprescindible para que la PWA se actualice."),
    ("h2", "Claude vía proxy con JWT"),
    ("p", "Anthropic nunca se llama desde el navegador: `api/anthropic.js` exige el JWT de Supabase (o `x-admin-key`) y aplica rate limit. Así la API key no se filtra y solo usuarios autenticados consumen cuota. El contrato con el modelo es JSON estricto, sin streaming."),
    ("h2", "Recomendador colaborativo entrenado offline"),
    ("p", "El SVD (scikit-surprise) corre en GitHub Actions cada día a las 3:00 UTC y materializa el top-20 por usuario en `recommendation_scores`. En producción no hay Python: `api/recommendations.js` solo lee filas con RLS. Es barato, desacoplado (cualquier otro algoritmo puede escribir en la misma tabla) y evita mantener un servicio de inferencia."),
    ("h2", "Migraciones SQL manuales"),
    ("p", "Las migraciones se corren a mano en el SQL Editor, en un orden documentado. Es simple para un equipo de una persona, pero no hay historial automático ni rollback, y los servicios toleran migraciones pendientes (degradación por columnas inexistentes) para poder desplegar el frontend antes. Adoptar la CLI de Supabase con migraciones versionadas es un paso natural cuando crezca el equipo."),
    ("h2", "Trueque: zonas fijas y matching por título"),
    ("p", "Para el MVP se eligieron 25 zonas fijas de CDMX (enum `exchange_zone`) en vez de geolocalización, y un match exacto sobre el título normalizado (columna GENERATED indexada) en vez de fuzzy matching o ISBN. Es determinista, rápido y explicable; a cambio se pierden coincidencias por ediciones o erratas. La migración a Google Places está en el roadmap."),
]

GUIDE = [
    ("h2", "Setup local paso a paso"),
    ("ol", [
        "Instala **Node.js 18 o superior** y Git (el proyecto usa Vite 5).",
        "Clona el repositorio y entra a la carpeta del proyecto.",
        "Instala dependencias: `npm install`.",
        "Crea tu archivo de entorno copiando la plantilla: `cp .env.example .env` (en Windows PowerShell: `Copy-Item .env.example .env`).",
        "Llena como mínimo `VITE_SUPABASE_URL`, `VITE_SUPABASE_ANON_KEY` y `ANTHROPIC_API_KEY` (ver Apéndice C).",
        "En un proyecto de Supabase propio, abre el SQL Editor y corre las migraciones **en el orden del Apéndice D** (columna «Orden de ejecución»). Desactiva «Confirm email» en Authentication → Providers → Email para que el registro deje sesión inmediata.",
        "Crea a mano en Storage los buckets públicos `avatars`, `covers` y `post-images` (el bucket `trueque` lo crea `trueque_schema.sql`).",
        "Arranca: `npm run dev` (levanta el proxy Anthropic local en el puerto 3001 y Vite en 5173 con `concurrently`).",
        "Abre http://localhost:5173, crea una cuenta y verifica que aparece en `public.users`.",
        "Opcional (recomendador): `pip install -r requirements.txt` y `python scripts/train_recommender.py --mock --dry-run`.",
        "Deploy a producción (manual): `vercel --prod`.",
    ]),
    ("h2", "Convenciones de código"),
    ("ul", [
        "**Idioma:** dominio, UI, comentarios y mensajes de error en español; identificadores técnicos en inglés cuando ya lo son (`fetchBooks`).",
        "**Servicios:** funciones async con *named exports* en `src/services/*Service.js`; devuelven datos o lanzan y el componente decide la UI de error.",
        "**Naming de datos:** columnas de la DB en `snake_case`; el mapeo a `camelCase` ocurre solo en los servicios (`dbToBook`/`bookToDb`). Excepción: el Trueque devuelve filas snake_case tal cual desde sus RPCs.",
        "**Componentes:** PascalCase; estilos inline con `palette`; modales como estado local + render condicional.",
        "**Recompensas:** jamás desde el cliente. Si algo otorga XP/gemas: trigger o RPC SECURITY DEFINER que llama a `folio_award` con un `ref` idempotente.",
        "**Errores de negocio en SQL:** `RAISE EXCEPTION 'PREFIJO_CODIGO'` y el servicio los traduce a un error tipado con mensaje en español (patrón `TruequeError`).",
        "**Fechas:** el cliente usa hora local (`localDateStr`), nunca UTC; el servidor calcula «hoy» en `America/Mexico_City`.",
        "**Al mover código fuera de App.jsx:** verifica exports/imports con grep antes de dar el trabajo por terminado.",
        "**Build limpio:** `npm run build` debe compilar sin errores antes de cada commit.",
    ]),
    ("h2", "Cómo agregar un feature nuevo (ejemplo: metas semanales de lectura)"),
    ("p", "Supón que quieres que cada usuario fije una meta de páginas por semana y vea su progreso. El proceso sigue siempre las mismas cinco capas."),
    ("h3", "Paso 1 — Tabla + RLS (nueva migración en supabase/)"),
    ("code", '''-- supabase/weekly_goals.sql
CREATE TABLE IF NOT EXISTS public.weekly_goals (
  user_id      uuid PRIMARY KEY REFERENCES public.users(id) ON DELETE CASCADE,
  pages_target int  NOT NULL CHECK (pages_target BETWEEN 10 AND 5000),
  updated_at   timestamptz NOT NULL DEFAULT now()
);
ALTER TABLE public.weekly_goals ENABLE ROW LEVEL SECURITY;

-- Patrón de auth_rls_migration.sql: lectura y escritura solo del dueño.
CREATE POLICY weekly_goals_own ON public.weekly_goals
  FOR ALL TO authenticated
  USING (auth.uid() = user_id) WITH CHECK (auth.uid() = user_id);'''),
    ("note", "Recuerda: con GRANTs por columna (como en `users`), una columna nueva puede requerir un `GRANT` explícito; y `permission denied for table X` es un problema de GRANT, no de policy."),
    ("h3", "Paso 2 — Servicio en src/services/"),
    ("code", '''// src/services/goalsService.js
import { supabase } from "../supabase.js";

export async function getWeeklyGoal(userId) {
  const { data, error } = await supabase
    .from("weekly_goals").select("pages_target").eq("user_id", userId).maybeSingle();
  if (error) throw error;
  return data?.pages_target ?? null;
}

export async function setWeeklyGoal(userId, pages) {
  const { error } = await supabase
    .from("weekly_goals").upsert({ user_id: userId, pages_target: pages });
  if (error) throw error;
}'''),
    ("h3", "Paso 3 — UI"),
    ("p", "Crea `src/components/WeeklyGoalCard.jsx` (props explícitas: `user`, `books`, `onChange`), usando `palette` para los colores, y móntalo desde `HomeView` o `MainApp`. Calcula el progreso en cliente con la fecha local (`localDateStr`) sumando `reading_logs` de la semana."),
    ("h3", "Paso 4 — Recompensa (solo si premia algo)"),
    ("p", "Si cumplir la meta da gemas, **no** lo hagas desde el cliente: crea una RPC SECURITY DEFINER que verifique en SQL que la meta se cumplió y llame a `folio_award(uid, xp, gems, 'weekly_goal', '<año-semana>')`. El `ref` idempotente evita cobrar dos veces la misma semana."),
    ("h3", "Paso 5 — Verificación y documentación"),
    ("ul", [
        "`npm run build` limpio y smoke test en el navegador (sin overlay rojo, sin errores en consola).",
        "Si hay recompensa: comprueba en Supabase que `reward_ledger` registró la fila **una sola vez**.",
        "Actualiza `DOCUMENTACION_ARQUITECTURA.md` (tablas, servicios, migración, changelog) y regenera el Word.",
    ]),
    ("h2", "Cómo debuggear los bugs más comunes"),
    ("table", ["Síntoma", "Causa probable", "Cómo verificar / arreglar"], [
        ["`permission denied for table users` (o cualquier tabla)", "Falta un **GRANT** de tabla/columna al rol `authenticated`. No es RLS (RLS dice «violates row-level security policy»). Con GRANTs por columna un UPSERT necesita más privilegios que un INSERT",
         "Consulta `information_schema.column_privileges`; corre `fix_signup_permissions.sql`; usa INSERT plano en vez de upsert"],
        ["`new row violates row-level security policy`", "Falta o no coincide una policy de INSERT/UPDATE (`auth.uid() = user_id`)", "`SELECT * FROM pg_policies WHERE tablename = '...'`; revisa que el cliente mande `user_id = auth.uid()`"],
        ["UPDATE/DELETE «no hace nada» sin error", "RLS filtra la fila: 0 filas afectadas", "Los servicios comparan filas afectadas (`booksService`); revisa la policy USING"],
        ["`42703` / `PGRST204` columna desconocida", "Migración sin correr en ese entorno", "Corre la migración; `isUnknownColumnError` degrada solo `total_pages/read_date_precision`"],
        ["`42P10` no unique constraint matching ON CONFLICT", "El target de `onConflict` no es una columna única", "Usa la PK o crea el UNIQUE; mejor INSERT plano + manejo de 23505"],
        ["`fetchStreakData is not a function` (o similar) tras mover código", "Función no exportada al extraerla del monolito", "grep del nombre; añade el `export`; envuelve en try/catch con fallback"],
        ["Pantalla roja en dev", "Overlay de debug (solo `import.meta.env.DEV`)", "Lee el stack trace; en producción solo se ve «Algo salió mal» + Recargar"],
        ["La PWA no se actualiza tras un deploy", "Cachea `index.html`/`sw.js`", "Revisa `vercel.json` (no-store) y desregistra el service worker en DevTools"],
        ["XP/gemas no suben", "Idempotencia (`reward_ledger`) o rate limit (12 logs/día, 4/hora)", "Busca la fila por `reason`+`ref`; `folio_award` ignora repetidos"],
        ["429 / respuesta sin `content` de Anthropic", "Rate limit del proxy (10 req/min/IP) o JWT vencido", "Reintenta; revisa `Authorization: Bearer`; ver mensajes amigables en `enrichBook`"],
        ["Cuenta privada no aparece en la búsqueda de amigos", "Efecto esperado de `privacy_hardening.sql`", "Diseñar una función/vista dedicada (`users_public`)"],
        ["`TRUEQUE_*` en la UI del Trueque", "Regla de negocio en SQL (`TRUEQUE_NO_ZONES`, `TRUEQUE_LIMIT_*`, `TRUEQUE_NO_CREDIT`…)", "`truequeService` los traduce a `TruequeError { code }`; ver Apéndice B"],
        ["No aparece ningún log `[auth]` en producción y el perfil sigue sin crearse", "Bundle viejo servido por el service worker, filtro de niveles en la consola, o la lógica no se ejecuta",
         "Abre la app en incógnito con DevTools y busca `[auth-debug] Build version`: debe traer `tag: 'auth-debug-1'` (`build` = hora real de compilación; `swControlled` = ¿lo sirve un service worker?). Luego sigue la traza `[auth-debug]` de `ensureUserProfile` (SELECT → INSERT payload → INSERT result)"],
        ["«Ya existe una cuenta con ese email» pero no se puede entrar", "Cuenta huérfana en `auth.users` sin perfil", "Iniciar sesión (se auto-repara) o reintentar el registro con la misma contraseña; en bloque: `repair_all_orphan_profiles.sql`"],
        ["Login OK pero `claim_daily_gems` o el onboarding de mascota fallan con foreign key / «usuario no existe»", "Cuenta huérfana: sin fila en `public.users`. Si el auto-perfil no se crea, mira la consola",
         "Busca `[auth]` en consola: `BLOQUEADO POR RLS`, `SIN PRIVILEGIOS` (corre `fix_signup_permissions.sql`) o el error completo; verifica que el frontend desplegado sea el nuevo (`vercel --prod`); repara en bloque con `repair_all_orphan_profiles.sql`"],
    ], [0.24, 0.36, 0.40]),
]

GLOSSARY = [
    ("Gemas", "Moneda virtual de FOLIO. Se ganan (libro terminado +50, sesión de lectura +5, bono diario, logros +10) y se gastan (BookTinder, búsqueda de matches del Trueque −20). Solo se leen en el cliente; toda escritura pasa por RPCs/triggers."),
    ("XP / Nivel", "Experiencia de la mascota. `xpForLevel = nivel × 100`, tope de nivel 50. Lo calcula el servidor en `folio_award`."),
    ("Mascota", "Compañero virtual del usuario (hoy un gato, «El Sensible») que sube de nivel con XP. Tabla `user_pets`; el cliente solo puede renombrarla."),
    ("Streak / Racha", "Días consecutivos leyendo. Filosofía anti-castigo: si falta un día la racha se **pausa**, no se resetea. Solo la avanza la RPC `update_streak()`."),
    ("Freeze / Protector", "Un protector de racha por mes (`streak_freezes_remaining`) que salva un día perdido. Se consume con `use_streak_freeze()`."),
    ("Trueque", "Módulo de intercambio de libros **físicos** entre usuarios de la misma zona de CDMX. Incluye matching, chat, cierre y calificación."),
    ("Match (perfecto / parcial)", "Empate del Trueque entre dos usuarios. Perfecto: cada uno tiene un libro que el otro busca. Parcial: solo una dirección. Uno vivo por pareja."),
    ("Crédito de búsqueda", "Vale de un solo uso (tabla interna `trueque_search_credits`) que emite `cost_search_matches` y consume `find_book_matches`, para que no se pueda buscar gratis."),
    ("Zona", "Una de las 25 zonas fijas de CDMX del Trueque (`exchange_zone`, `src/config/zones.js`). Free elige 2–3, Plus 2–5."),
    ("Folio Plus (`is_premium`)", "Plan de pago del Trueque con límites mayores. Aún sin pasarela: se activa a mano en SQL."),
    ("`folio_award`", "Función SECURITY DEFINER, única puerta de otorgamiento de XP/gemas. Idempotente vía `reward_ledger`; revocada a clientes."),
    ("`reward_ledger`", "Libro mayor de recompensas: una fila por (usuario, motivo, ref) única. Garantiza que un premio no se pague dos veces."),
    ("Server authority", "Principio: todo valor «ganable» se decide en Postgres, no en el cliente."),
    ("RLS", "Row Level Security de Postgres: policies por fila (`auth.uid() = user_id`). Complementa —no sustituye— a los GRANTs de tabla/columna."),
    ("SECURITY DEFINER", "Modo de función que se ejecuta con los privilegios de su dueño; permite escribir en tablas que el cliente no puede tocar."),
    ("RPC", "Llamada a una función de Postgres desde el cliente con `supabase.rpc('nombre', { ... })`."),
    ("`session_id`", "UUID generado por el cliente por sesión de lectura; sirve de `ref` idempotente y sobrevive a reintentos offline."),
    ("Cuenta huérfana", "Usuario que existe en `auth.users` pero no tiene fila en `public.users` (signup que falló a medias). Se auto-repara en login, sesión guardada y `onAuthStateChange`; el histórico, con `repair_all_orphan_profiles.sql`."),
    ("`ensureUserProfile`", "Función de `authService` que comprueba si existe el perfil del usuario y, si falta, lo crea (idempotente, con reintento de username y logs de error explícitos)."),
    ("`[auth-debug]`", "Prefijo de los logs de diagnóstico de auth (temporales). Se apagan con `AUTH_DEBUG = false` en `authService.js`. `[auth-debug] Build version` (en `main.jsx`) indica qué build corre en el navegador."),
    ("`watchAuthProfile`", "Suscripción a `onAuthStateChange` (`SIGNED_IN` / `INITIAL_SESSION`) que ejecuta `ensureUserProfile`; la monta `App`."),
    ("BookTinder", "Swipe de recomendaciones sobre `books_curated`, límite de 15 guardados/día ampliable con gemas."),
    ("Snacks", "Cuentos de dominio público (~23) embebidos en la app y leídos en `ReaderView`, sin red."),
    ("Wrapped", "Resumen mensual estilo Spotify (libros, páginas, racha), exportable como imagen."),
    ("Logro (achievement)", "Hito desbloqueable (27 definidos). El cliente propone candidatos y `award_achievement(key)` los valida en SQL."),
    ("Event bus", "Pub/sub casero (`petBus`, `gemsEventBus`, `achievementBus`) para que servicios sin React notifiquen a la UI."),
    ("`users_public`", "Vista `security_invoker` con columnas públicas de `users` (sin email). Creada por `privacy_hardening.sql`; aún sin uso."),
    ("Embajadas FOLIO", "Lugares aliados verificados para intercambios (pendiente de Sprint 2; el disclaimer ya las anuncia)."),
    ("Buscalibre / afiliado", "Tienda a la que apuntan los links de compra, con `VITE_BUSCALIBRE_AFFILIATE_ID`."),
    ("UAM", "Universidad Autónoma Metropolitana; su catálogo curado vive en `UAMLibraryView` y los libros llevan `is_uam_book`."),
    ("PWA", "Progressive Web App: la app instalable con service worker (Workbox) y soporte offline."),
    ("Monolito App.jsx", "El archivo de ~14.900 líneas que contiene casi todas las vistas."),
    ("INDAUTOR", "Instituto Nacional del Derecho de Autor (México). El sprint de seguridad del 2026-09-17 se nombró «pre-INDAUTOR» (por el nombre del commit)."),
]

# Contexto adicional por sección del .md (clave = número romano)
EXTRA = {
    "I": [
        ("p", "Esta sección resume el «qué» y el «con qué». Para entender el sistema conviene tener en la cabeza una idea: **el navegador es un cliente delgado y la base de datos es la autoridad**. Casi toda la lógica sensible vive en Postgres y las funciones de Vercel son solo tuberías con credenciales."),
        ("code", '''Navegador (PWA, React)
   │  supabase-js (JWT del usuario)                 │  fetch + Bearer JWT
   ▼                                                ▼
Supabase                                     Vercel /api/*
 ├─ Auth (email + password)                   ├─ anthropic.js  → api.anthropic.com
 ├─ Postgres + RLS                            ├─ books.js      → Google Books
 │   └─ triggers / RPCs SECURITY DEFINER      └─ recommendations.js → lee scores con RLS
 └─ Storage (avatars, covers, post-images, trueque)
GitHub Actions (cron 3:00 UTC) → train_recommender.py → recommendation_scores'''),
        ("p", "**Ejemplo de lectura del diagrama:** cuando un usuario termina un libro, el cliente solo hace `UPDATE books SET status='read'`. Un trigger en Postgres detecta el cambio y llama a `folio_award`, que suma XP y gemas y hace el level-up de la mascota. El cliente después *lee* el resultado; nunca lo escribe."),
    ],
    "II": [
        ("p", "Regla práctica para orientarse: `src/services/` es la capa de datos, `src/components/` y `src/pets/` contienen lo extraído del monolito, `src/App.jsx` contiene todo lo demás, y `supabase/` son las migraciones. Todo lo que empiece con `api/` corre en Vercel."),
        ("ul", [
            "**¿Dónde busco un bug de datos?** Primero en `src/services/`, luego en la RPC/trigger correspondiente en `supabase/`.",
            "**¿Dónde busco una pantalla?** `grep -n \"function NombreVista\" src/App.jsx`; los comentarios `// ============ SECCIÓN ============` son el índice.",
            "**¿Dónde configuro el caché de la PWA?** `vite.config.js` (Workbox) y `vercel.json` (headers).",
        ]),
    ],
    "III": [
        ("p", "El stack está pensado para que el equipo pueda ser de una persona: no hay servidores que operar, ni Docker, ni CI complejo. La única pieza «pesada» es el entrenamiento del recomendador, aislado en GitHub Actions."),
        ("code", '''Flujo de una llamada a Claude (mood / enriquecer libro)
cliente ── POST /api/anthropic  (Authorization: Bearer <JWT Supabase>) ──►
   api/anthropic.js:  supabase.auth.getUser(jwt)  → 401 si inválido
                      rate limit 10 req/min/IP (Map en memoria)
                      ──► api.anthropic.com/v1/messages  (ANTHROPIC_API_KEY)
cliente ◄── JSON  → valida 429 / payload sin `content` → mensaje amigable'''),
        ("note", "**Variables de entorno:** las `VITE_*` se incrustan en el bundle (son públicas por diseño); todo lo demás es solo de servidor. `SUPABASE_SERVICE_ROLE_KEY` salta RLS y solo debe existir como secret de GitHub Actions. Detalle completo en el Apéndice C."),
    ],
    "IV": [
        ("p", "Cada feature sigue el mismo esqueleto: **vista/modal en React → servicio → tabla con RLS → (si premia algo) trigger o RPC SECURITY DEFINER**. Los ejemplos siguientes muestran ese esqueleto en las dos features más representativas."),
        ("code", '''Ejemplo — Racha
ReadingLogModal ─► INSERT reading_logs {pages_read, mood, log_date, session_id}
                      └─ trigger trg_reading_logged ─► folio_award(+XP, +gemas)
                ─► RPC update_streak()  (exige un log real de hoy, fecha México)
UI ◄─ fetchStreakData(userId) → { streak, hasLoggedToday, pagesLoggedToday }'''),
        ("code", '''Ejemplo — Trueque (resumen)
ZonesOnboarding ─► set_exchange_zones
OfferBookModal  ─► INSERT books_offered      WantBookModal ─► INSERT books_wanted
«Buscar matches» ─► cost_search_matches (−20 💎 o cupo Plus) ─► find_book_matches
MatchCard ─► LegalDisclaimerModal ─► ExchangeChat ─► complete_exchange ─► rate_exchange'''),
    ],
    "V": [
        ("p", "Los flujos críticos son los que cruzan más capas. Tres invariantes ayudan a razonar sobre ellos: (1) el `id` de `public.users` es siempre el `auth.uid()`; (2) todo premio es idempotente; (3) las escrituras offline se encolan y reintentan con el mismo identificador."),
        ("h3", "Diagramas de auth: registro y auto-reparación de cuentas huérfanas"),
        ("code", '''REGISTRO
signUp OK ──► INSERT perfil OK ──────────────► onboarding ✔
   │             │
   │             └─ falla (p. ej. permisos) ──► signOut ──► cuenta HUÉRFANA en auth.users
   │
   └─ «ya registrado» ──► signIn con la misma contraseña
                              ├─ perfil existe ──► «Ya existe una cuenta»
                              └─ perfil NO existe ──► INSERT perfil ──► entra ✔ (recuperada)

AUTO-REPARACIÓN (2026-09-19) — se ejecuta en TODOS los caminos que dejan sesión
login ───────────────► buildAppUser ─┐
sesión guardada ─────► buildAppUser ─┼─► ensureUserProfile(authUser)   (una sola promesa por usuario)
onAuthStateChange ───► setTimeout(0)─┘        │
 (SIGNED_IN | INITIAL_SESSION; se ignora       ├─ SELECT falla ─► console.error, NO crea; usa caché o null
  mientras corre el registro)                  ├─ hay fila ─────► ok
                                               └─ no hay fila ──► INSERT {id, email, nombre, username,
                                                                          is_public:true, onboarding_completed:false}
                                                    ├─ username duplicado ─► sufijo aleatorio, máx. 3 reintentos
                                                    ├─ PK duplicada ───────► ok (creado en paralelo)
                                                    └─ RLS / GRANT / otro ─► console.error específico;
                                                                             login degradado (profileMissing)'''),
        ("note", "**Por qué en todos los caminos:** con un solo punto de reparación (el alta) una cuenta huérfana con sesión guardada nunca se arreglaba, y el onboarding de mascota y `claim_daily_gems` fallaban por FK contra `users`. El histórico se repara en bloque con `supabase/repair_all_orphan_profiles.sql`."),
    ],
    "VI": [
        ("p", "La capa de servicios nació de extraer código del monolito (commit `3b3adfd`). Su contrato: funciones async con named exports que **devuelven datos o lanzan**; no muestran UI. Los helpers que aún están en App.jsx (`logReadingSession`, `checkAchievements`, `enrichBook`, `searchGoogleBooks`, `createFeedPost`…) son candidatos naturales a moverse aquí."),
        ("code", '''// Patrón típico de un servicio con caché offline
export async function fetchBooks(userId) {
  const { data, error } = await supabase.from("books").select("*").eq("user_id", userId);
  if (error) return getCachedBooks(userId);   // cae a localStorage
  const books = data.map(dbToBook);           // snake_case → camelCase
  cacheBooks(userId, books);
  return books;
}'''),
    ],
    "VII": [
        ("p", "Todo componente recibe `user`, `books`, `setTab` y callbacks por props desde `MainApp`; los modales son estado local con render condicional. Para localizar un componente en App.jsx no hay índice: se busca por nombre con grep."),
        ("code", '''MainApp (≈40 estados)
 ├─ header + BottomNav (6 tabs: home, social, library/add, trueque, profile…)
 ├─ vista activa (HomeView | SocialView | LibraryView | AddBookView | TruequeMain | PerfilWrapper)
 └─ overlays globales (PetHub, NotificationsSheet, celebraciones, toasts)'''),
    ],
    "VIII": [
        ("p", "El modelo de seguridad de la base tiene **tres capas** que hay que entender juntas: (1) **GRANTs** de tabla y de columna al rol `authenticated`; (2) **policies RLS** por fila; (3) **triggers/RPCs SECURITY DEFINER** para lo que el cliente nunca debe escribir. Un acceso solo prospera si pasa las tres."),
        ("code", '''Petición del cliente ──► ¿GRANT de tabla/columna?  ── no ──► permission denied for table X
                              │ sí
                              ▼
                        ¿policy RLS de la fila?   ── no ──► row-level security policy violation / 0 filas
                              │ sí
                              ▼
                        ¿triggers BEFORE/AFTER?   ── p. ej. guard_users_premium, guard_pet_columns
                              ▼
                          fila escrita'''),
        ("note", "**Caso de estudio (signup, septiembre 2026):** `privacy_hardening.sql` revocó SELECT sobre `users` y lo re-otorgó por columna (sin `email`). El registro usaba un UPSERT, que necesita más privilegios que un INSERT; el error resultante fue `permission denied for table users` (capa 1), no una violación de RLS (capa 2). La corrección fue doble: INSERT plano en el cliente y `fix_signup_permissions.sql` con policies explícitas más GRANT INSERT/UPDATE por columna. Ver también los patrones 12 y 13 (cuentas huérfanas y auto-reparación) en el capítulo de patrones."),
    ],
    "IX": [
        ("p", "Ninguna integración externa se llama con secretos desde el navegador: o pasan por `api/*` (Anthropic, Google Books) o usan la anon key con RLS (Supabase). Las únicas llamadas directas son públicas y de solo lectura (Open Library para portadas)."),
    ],
    "X": [
        ("p", "Los patrones son las «reglas del juego» del código. Si vas a romper alguno, tiene que ser a propósito y quedar documentado. Los más importantes son el 5 (idempotencia), el 6 (server authority) y el 13 (auto-reparación de perfiles huérfanos)."),
        ("code", '''// Patrón 2 — event bus
const unsub = petBus.on((evt) => setPetToast(evt));   // dentro de useEffect
return () => unsub();

// Patrón 5 — idempotencia (SQL)
INSERT INTO reward_ledger (user_id, reason, ref) VALUES (uid, 'book_finished', book_id)
ON CONFLICT DO NOTHING;   -- si ya existía, folio_award no vuelve a pagar'''),
    ],
    "XI": [
        ("p", "Los puntos de extensión están pensados para crecer sin tocar el núcleo. La receta general está desarrollada con un ejemplo completo en el capítulo «Guía para nuevos developers»."),
    ],
    "XII": [
        ("p", "La deuda técnica se lista con honestidad porque condiciona el riesgo de cambiar el sistema. Regla de lectura: lo **Pendiente** es trabajo abierto; lo **Resuelto** documenta cómo se cerró para que no se reintroduzca (por ejemplo, no volver a escribir logros o racha desde el cliente)."),
    ],
    "XIII": [
        ("p", "Este checklist es el mínimo para ser productivo en un día. Complétalo con la Guía para nuevos developers (setup, convenciones y debugging) de este documento."),
    ],
}

# ---------------------------------------------------------------------------
# Apéndices (datos curados; A y B se contrastan con supabase/*.sql)
# ---------------------------------------------------------------------------
CORE_SCHEMA = {
    "users": ("Perfil del usuario; `id = auth.uid()`.",
              [("id", "uuid, PK (= auth.uid())"), ("nombre", "text (nombre visible; el cliente lo mapea a `name`)"), ("username", "text (asumido único)"),
               ("email", "text (SIN SELECT para clientes)"), ("avatar_url", "text"), ("cover_url", "text"), ("bio", "text"),
               ("is_public", "boolean, default true"), ("preferred_genres", "text[] (onboarding)"), ("onboarding_completed", "boolean"),
               ("is_premium", "boolean NOT NULL DEFAULT false (Folio Plus; protegido por `guard_users_premium`)")]),
    "books": ("Biblioteca por usuario (por valor, sin FK a catálogo).",
              [("id", "uuid PK"), ("user_id", "uuid → users"), ("title / author", "text"), ("status", "reading | want_to_read | wish | read"),
               ("genre / summary / review", "text"), ("rating", "int"), ("cover_url", "text"), ("mood_tags", "text[]"), ("added_at / finished_at", "timestamptz"),
               ("read_date_precision", "text, default 'exact'"), ("is_uam_book", "boolean"), ("isbn", "text"), ("total_pages", "int")]),
    "quotes": ("Citas guardadas.", [("id", "uuid PK"), ("user_id", "uuid"), ("book_id", "uuid → books"), ("text", "text"), ("page_number", "int"), ("mood", "text (legado)"), ("tags", "text[]"), ("is_favorite", "boolean"), ("is_public", "boolean"), ("created_at", "timestamptz")]),
    "reading_logs": ("Sesiones de lectura.", [("user_id", "uuid"), ("book_id", "uuid"), ("pages_read", "int"), ("mood", "text"), ("log_date", "date (local del cliente)"), ("session_id", "uuid NOT NULL DEFAULT gen_random_uuid(), UNIQUE")]),
    "user_streaks": ("Racha (1:1 con el usuario; solo SELECT para el cliente).", [("user_id", "uuid"), ("current_streak", "int"), ("longest_streak", "int"), ("last_log_date", "date"), ("total_pages_read", "int"), ("streak_freeze_used_at", "date"), ("streak_freezes_remaining", "int"), ("updated_at", "timestamptz")]),
    "user_gems": ("Saldo de gemas (solo SELECT propio).", [("user_id", "uuid"), ("balance", "int"), ("(columnas de control del bono diario)", "ver `claim_daily_gems` en sprint1_server_authority.sql")]),
    "user_pets": ("Mascota (1:1; el cliente solo renombra).", [("id", "uuid PK"), ("user_id", "uuid UNIQUE NOT NULL → users ON DELETE CASCADE"), ("pet_type", "text"), ("pet_name", "text"), ("xp", "int"), ("level", "int"), ("created_at / updated_at", "timestamptz")]),
    "achievements": ("Logros desbloqueados (solo escribe `award_achievement`).", [("user_id", "uuid"), ("key", "text"), ("(UNIQUE user_id, key)", "un logro por usuario")]),
    "monthly_wraps": ("Wrapped mensual.", [("user_id", "uuid"), ("year / month", "int"), ("wrap_type", "text"), ("(payload)", "contenido del resumen")]),
    "notifications": ("Notificaciones (el actor inserta; lee el destinatario).", [("user_id", "uuid (destinatario)"), ("actor_id", "uuid"), ("type", "text"), ("related_id", "uuid"), ("read", "boolean")]),
    "friendships": ("Amistades.", [("user_id", "uuid"), ("friend_id", "uuid"), ("status", "pending | accepted")]),
    "conversations": ("Chats 1:1.", [("id", "uuid PK"), ("user1_id / user2_id", "uuid")]),
    "messages": ("Mensajes de chat.", [("conversation_id", "uuid"), ("sender_id", "uuid"), ("(texto y created_at)", "ver ChatView")]),
    "posts": ("Posts del feed.", [("id", "uuid PK"), ("user_id", "uuid"), ("type", "tipo de post (sesión, logro, cita, libro terminado)"), ("content", "contenido del post")]),
    "comments": ("Comentarios de posts.", [("post_id", "uuid"), ("user_id", "uuid"), ("(texto)", "text")]),
    "comment_replies": ("Respuestas a comentarios.", [("comment_id", "uuid"), ("user_id", "uuid"), ("(texto)", "text")]),
    "post_likes": ("Likes de posts.", [("post_id", "uuid"), ("user_id", "uuid")]),
    "comment_likes": ("Likes de comentarios.", [("comment_id", "uuid"), ("user_id", "uuid")]),
    "books_curated": ("Catálogo de ~500 libros (seed en scripts/).", [("(title, author, género, rating…)", "ver scripts/books-500.sql")]),
    "daily_save_limits": ("Límite diario de BookTinder (client-side).", [("user_id", "uuid"), ("date", "date"), ("saves_used", "int"), ("gems_spent", "int")]),
}

RPCS = [
    ("pet_daily_checkin()", "sprint1 / timezone_fix", "+3 XP si la racha está activa; idempotente por día (México).", "await supabase.rpc('pet_daily_checkin');"),
    ("claim_daily_gems()", "sprint1 / timezone_fix", "Bono diario (+5 y consecutivos), gate de 20 h; crea la fila de gemas con bienvenida.", "const { data } = await supabase.rpc('claim_daily_gems');"),
    ("claim_achievement_gems(p_keys text[])", "sprint1", "+10 gemas por logro, idempotente por key. Devuelve el saldo.", "await supabase.rpc('claim_achievement_gems', { p_keys: ['first_book'] });"),
    ("award_achievement(p_key text) → boolean", "security_patch_achievements", "Revalida en SQL la condición real y otorga el logro; `false` si ya estaba o no cumple.", "const { data: ok } = await supabase.rpc('award_achievement', { p_key: 'first_book' });"),
    ("update_streak()", "security_patch_streaks", "Avanza racha y total de páginas; exige un `reading_log` real de hoy.", "await supabase.rpc('update_streak');"),
    ("reset_monthly_freeze()", "security_patch_streaks", "Reasigna 1 protector si cambió el mes y quedaban 0.", "await supabase.rpc('reset_monthly_freeze');"),
    ("use_streak_freeze()", "security_patch_streaks", "Consume el protector; lanza excepción si no hay.", "await supabase.rpc('use_streak_freeze');"),
    ("delete_my_account()", "features_migration", "Borra en cascada la cuenta del usuario autenticado.", "await supabase.rpc('delete_my_account');"),
    ("trueque_status() → jsonb", "trueque_schema", "Plan, límites efectivos, costo de búsqueda, búsquedas de hoy y conteos.", "const { data } = await supabase.rpc('trueque_status');"),
    ("set_exchange_zones(p_zones text[])", "trueque_schema", "Reemplazo atómico de zonas; valida 2..máx del plan.", "await supabase.rpc('set_exchange_zones', { p_zones: ['roma', 'condesa'] });"),
    ("cost_search_matches(p_user_id uuid) → boolean", "trueque_schema", "Cobra 20 gemas (o cupo Plus) y emite un crédito; `false` si no alcanza.", "const { data: paid } = await supabase.rpc('cost_search_matches', { p_user_id: user.id });"),
    ("find_book_matches(p_user_id uuid) → TABLE", "trueque_schema", "Consume el crédito, genera/actualiza matches y devuelve la lista.", "const { data } = await supabase.rpc('find_book_matches', { p_user_id: user.id });"),
    ("get_my_exchange_matches() → TABLE", "trueque_schema", "Lista los matches sin descubrir ni cobrar.", "const { data } = await supabase.rpc('get_my_exchange_matches');"),
    ("send_exchange_message(p_match_id uuid, p_message text) → uuid", "trueque_schema", "Envía un mensaje al chat; pending→chatting.", "await supabase.rpc('send_exchange_message', { p_match_id: id, p_message: 'Hola' });"),
    ("complete_exchange(p_match_id uuid) → boolean", "trueque_schema", "Marca el intercambio como completado y desactiva los libros.", "await supabase.rpc('complete_exchange', { p_match_id: id });"),
    ("rate_exchange(p_match_id uuid, p_rating int, p_review text) → uuid", "trueque_schema", "Califica 1–5 (una vez por participante).", "await supabase.rpc('rate_exchange', { p_match_id: id, p_rating: 5, p_review: 'Excelente' });"),
]
RPC_INTERNAL = [
    ("folio_award(user, xp, gems, reason, ref)", "sprint1", "Única puerta de recompensas. REVOKE a clientes: solo triggers/RPCs."),
    ("cleanup_expired_matches() → int", "trueque_schema", "pg_cron diario (`0 9 * * *` UTC). No expuesta a clientes."),
    ("trueque_norm / trueque_cfg / trueque_generate_matches / trueque_matches_for", "trueque_schema", "Auxiliares internas del Trueque (REVOKE a clientes)."),
]

ENV_VARS = [
    ("VITE_SUPABASE_URL", "Cliente + servidor + Actions", "Sí", "URL del proyecto Supabase. Pública; la usan el cliente, `api/*`, `server.js` y el job de Python."),
    ("VITE_SUPABASE_ANON_KEY", "Cliente + servidor + Actions", "Sí", "Key anon/public. Pública; la seguridad la da RLS. Las funciones `api/*` la usan con el JWT del usuario."),
    ("VITE_BUSCALIBRE_AFFILIATE_ID", "Cliente", "No", "ID de afiliado de Buscalibre; App.jsx trae un valor por defecto."),
    ("ANTHROPIC_API_KEY", "Solo servidor", "Sí (para IA)", "Clave del proxy `/api/anthropic` y `server.js`. Sin prefijo VITE_: jamás al navegador."),
    ("ADMIN_KEY", "Solo servidor", "No", "Bypass del proxy Anthropic con el header `x-admin-key` (sin JWT ni rate limit). Usar un valor largo y aleatorio."),
    ("GOOGLE_BOOKS_API_KEY", "Solo servidor", "No", "Sube la cuota de `/api/books`; funciona sin ella."),
    ("SUPABASE_SERVICE_ROLE_KEY", "Solo GitHub Actions (secret)", "Solo para el job", "Salta RLS. Solo para entrenar el recomendador y seeds. Jamás en Vercel ni con prefijo VITE_."),
    ("SUPABASE_SECRET_KEY", "Solo `scripts/generate-500-books.js`", "No", "Escritura del catálogo curado."),
    ("PORT", "Solo `server.js` (dev)", "No", "Puerto del proxy local. Por defecto 3001."),
    ("SUPABASE_URL / SUPABASE_ANON_KEY", "Job de Python", "No", "Alternativa a las `VITE_*` para `train_recommender.py`."),
]

MIGRATIONS = [
    # (orden de ejecución, archivo, fecha fallback, qué hace)
    (1, "auth_rls_migration.sql", "2026-07-02", "Migración a Supabase Auth + RLS de todas las tablas core y policies de Storage. Incluye un WIPE destructivo (bloque 0)."),
    (2, "features_migration.sql", "2026-07-15", "`user_lists`/`user_list_books`, `delete_my_account()`, `is_public`, columnas de lectura."),
    (3, "recommendation_scores.sql", "2026-07-02", "Tabla que llena el job de Python."),
    (4, "sprint1_server_authority.sql", "2026-07-16", "`reward_ledger`, `folio_award`, triggers anti-cheat, RPCs de XP/gemas."),
    (5, "security_patch_reading_logs.sql", "2026-09-17", "`session_id` único y rate limit de recompensas de lectura."),
    (6, "security_patch_achievements.sql", "2026-09-17", "Cierra el INSERT libre de logros → RPC `award_achievement`."),
    (7, "security_patch_streaks.sql", "2026-09-17", "Cierra el UPDATE libre de `user_streaks` → RPCs de racha/freeze."),
    (8, "privacy_hardening.sql", "2026-09-17", "RLS de lectura por dueño/público/amigo; `users.email` fuera del SELECT (REVOKE + GRANT por columna)."),
    (9, "timezone_fix.sql", "2026-09-17", "«Hoy» server-side = America/Mexico_City."),
    (10, "trueque_schema.sql", "2026-09-19", "Trueque: 7 tablas + RLS + RPCs + bucket + pg_cron; `users.is_premium`."),
    (11, "fix_signup_permissions.sql", "2026-09-18", "Fix del signup: policies INSERT/UPDATE de `users` y GRANT INSERT/UPDATE por columna."),
    ("—", "cleanup_orphan_auth_users.sql", "2026-09-18", "Opcional: lista, repara o borra cuentas huérfanas de `auth.users`."),
    (12, "repair_all_orphan_profiles.sql", "2026-09-19", "Repara en bloque TODAS las cuentas huérfanas (perfil desde `raw_user_meta_data`, usernames duplicados con sufijo). Idempotente."),
]


# ============================================================================
# 6. Ensamblado
# ============================================================================
def appendix_a(b):
    b.begin_appendix("A", "Esquema completo de tablas")
    b.paragraph("Las tablas del **Apartado A.1** se leen automáticamente de `supabase/*.sql` (fuente fiable). Las del **A.2** se crearon "
                "desde el dashboard de Supabase y no tienen DDL en el repo: sus columnas se reconstruyeron del código y de la auditoría "
                "de fase 1, por lo que pueden faltar columnas de control (`id`, `created_at`…). Verifica siempre contra la base real con "
                "`SELECT column_name, data_type FROM information_schema.columns WHERE table_name = '<tabla>'`.")
    b.heading(2, "Tablas definidas en migraciones SQL")
    sql_tables = parse_sql_tables()
    if not sql_tables:
        b.note("No se encontró `supabase/*.sql` (¿se ejecutó el script fuera del repo?).")
    for name, (fname, cols) in sql_tables.items():
        b.heading(3, f"`{name}`  —  {fname}")
        b.table(["Columna", "Definición"], [[f"`{c}`" if not c.startswith("(") else c, d] for c, d in cols], [0.24, 0.76])
    b.heading(2, "Tablas creadas desde el dashboard (reconstruidas del código)")
    for name, (desc, cols) in CORE_SCHEMA.items():
        b.heading(3, f"`{name}`")
        b.paragraph(desc)
        b.table(["Columna", "Tipo / notas"], [[f"`{c}`" if not c.startswith("(") else c, t] for c, t in cols], [0.32, 0.68])
    b.paragraph("**Vista:** `users_public` (`security_invoker`): `id, nombre, username, avatar_url, cover_url, bio, is_public`. "
                "**Storage:** buckets `avatars`, `covers`, `post-images` (manuales) y `trueque` (creado por SQL).")


def appendix_b(b):
    b.begin_appendix("B", "RPCs: firma y ejemplo de uso")
    known = parse_sql_functions()
    missing = [n for n in (re.match(r"(\w+)", r[0]).group(1) for r in RPCS) if known and n not in known]
    if missing:
        print(f"AVISO: RPCs del apéndice B no encontradas en supabase/*.sql: {', '.join(missing)}", file=sys.stderr)
    b.paragraph("Todas se llaman con `supabase.rpc(nombre, args)` desde un usuario autenticado y son SECURITY DEFINER. "
                "Los errores de negocio llegan como `RAISE EXCEPTION 'CODIGO'`.")
    b.heading(2, "RPCs expuestas al cliente")
    for sig, src, desc, ex in RPCS:
        b.heading(3, f"`{sig}`")
        b.paragraph(f"{desc} (definida en `{src}`)")
        b.code(ex)
    b.heading(2, "Funciones internas (no llamables desde el cliente)")
    b.table(["Función", "Definida en", "Qué hace"], [[f"`{a}`", s, d] for a, s, d in RPC_INTERNAL], [0.4, 0.2, 0.4])
    b.heading(2, "Códigos de error del Trueque")
    b.table(["Código", "Cuándo ocurre"], [
        ["`TRUEQUE_NO_ZONES` / `TRUEQUE_MIN_ZONES`", "Quien busca no tiene zonas / se eligieron menos del mínimo"],
        ["`TRUEQUE_LIMIT_ZONES` / `_OFFERED` / `_WANTED`", "Límite del plan free/Plus (la UI abre el modal de Plus)"],
        ["`TRUEQUE_NO_GEMS`", "Gemas insuficientes para la acción"],
        ["`TRUEQUE_NO_CREDIT`", "`find_book_matches` sin un crédito emitido por `cost_search_matches`"],
        ["`TRUEQUE_RATE_LIMIT`", "Tope diario de búsquedas de Plus alcanzado"],
        ["`TRUEQUE_NOT_PARTICIPANT`", "El usuario no es parte del match"],
        ["`TRUEQUE_MATCH_CLOSED` / `_MATCH_EXPIRED`", "El match ya no está vivo / venció"],
        ["`TRUEQUE_BAD_MESSAGE` / `_BAD_RATING`", "Mensaje vacío o demasiado largo / calificación fuera de 1–5"],
        ["`TRUEQUE_NOT_COMPLETED` / `_ALREADY_RATED`", "Calificar antes de completar / ya calificó"],
        ["`TRUEQUE_BOOK_IN_MATCH`", "No se puede cambiar título/autor de un libro en un match vivo"],
    ], [0.4, 0.6])


def appendix_c(b):
    b.begin_appendix("C", "Variables de entorno")
    b.paragraph("Las variables con prefijo `VITE_` se incrustan en el bundle del navegador; el resto es solo de servidor. "
                "La plantilla comentada está en `.env.example`.")
    b.table(["Variable", "Dónde", "Obligatoria", "Explicación"], [[f"`{a}`", w, o, e] for a, w, o, e in ENV_VARS], [0.25, 0.2, 0.12, 0.43])


def appendix_d(b):
    b.begin_appendix("D", "Historial de migraciones SQL")
    b.paragraph("Ordenadas **cronológicamente** por la fecha del commit que las introdujo (git). El «orden de ejecución» es el que debe seguirse "
                "en un proyecto Supabase nuevo (respeta dependencias).")
    rows = []
    for order, fname, fallback, desc in MIGRATIONS:
        p = SQL_DIR / fname
        date = git_first_date(p) if p.exists() else ""
        rows.append((date or fallback, order, fname, desc))
    rows.sort(key=lambda r: (r[0], str(r[1])))
    b.table(["Fecha", "Orden de ejecución", "Archivo", "Qué hace"],
            [[d, str(o), f"`{f}`", desc] for d, o, f, desc in rows], [0.12, 0.13, 0.3, 0.45])
    b.note("La fecha es la del commit que añadió el archivo. `fix_signup_permissions.sql`, `cleanup_orphan_auth_users.sql` y `repair_all_orphan_profiles.sql` usan la fecha del "
           "changelog hasta que se comiteen. Antes de una migración destructiva (`auth_rls_migration.sql` bloque 0) haz respaldo.")


def build(doc, toc_entries):
    setup_styles(doc)
    md = MD_PATH.read_text(encoding="utf-8")
    pre, sections = split_sections(md)
    b = Builder(doc, toc_entries)
    b.cover()
    b.toc()

    # 1. Introducción (+ changelog del documento fuente)
    b.heading(1, "Introducción")
    b.render(INTRO)
    changelog = next((s for s in sections if s["roman"] is None and s["name"].upper() == "CHANGELOG"), None)
    if changelog:
        b.heading(2, "Historial de cambios del documento fuente")
        render_md_blocks(b, parse_blocks(changelog["lines"]))
    updated = next((ln for ln in pre if "Última actualización" in ln), "")
    if updated:
        b.note(updated.lstrip("> ").strip())

    # 2..14. Secciones del .md con contexto adicional
    for s in sections:
        if s["roman"] is None:
            continue
        b.heading(1, nice_title(s["name"]))
        extra = EXTRA.get(s["roman"])
        if extra:
            b.heading(2, "Contexto adicional")
            b.render(extra)
        render_md_blocks(b, parse_blocks(s["lines"]))

    # Capítulos nuevos
    b.heading(1, "Decisiones de arquitectura y por qué")
    b.render(DECISIONS)
    b.heading(1, "Guía para nuevos developers")
    b.render(GUIDE)
    b.heading(1, "Glosario")
    b.paragraph("Términos específicos de FOLIO, en orden de aparición temática.")
    b.table(["Término", "Significado"], [[t if "`" in t else f"**{t}**", d] for t, d in GLOSSARY], [0.24, 0.76])

    # Apéndices
    appendix_a(b)
    appendix_b(b)
    appendix_c(b)
    appendix_d(b)

    # Historial reciente de git (solo si hay repo)
    log = git("log", "-15", "--format=%ad|%h|%s", "--date=short")
    if log:
        rows = [ln.split("|", 2) for ln in log.splitlines()]
        if PENDING_COMMIT:   # el .docx se genera ANTES de comitearse: refleja el commit en curso
            rows.insert(0, [datetime.date.today().isoformat(), "(este commit)", PENDING_COMMIT])
            rows = rows[:15]
        b.app = None
        b.begin_appendix("E", "Últimos commits del repositorio")
        b.table(["Fecha", "Commit", "Mensaje"], rows, [0.13, 0.12, 0.75])
    return b.headings


# Términos de auth que deben decir lo mismo en el .md y en el .docx (auditoría de coherencia)
AUTH_TERMS = ["ensureUserProfile", "watchAuthProfile", "repair_all_orphan_profiles.sql", "fix_signup_permissions.sql",
              "onAuthStateChange", "recoverOrphanAccount", "INITIAL_SESSION", "[auth-debug]", "AUTH_DEBUG", "__APP_BUILD__"]


def all_docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def verify_sync():
    """Comprueba que el .docx refleje el .md: encabezados, migraciones SQL y términos de auth. Devuelve nº de problemas."""
    md = MD_PATH.read_text(encoding="utf-8")
    text = all_docx_text(OUT_PATH)
    flat = re.sub(r"\s+", " ", text.replace("`", ""))
    problems = []
    for m in re.finditer(r"^#{2,4}\s+(.*)$", md, re.M):
        raw = m.group(1).strip()
        if raw.upper() == "CHANGELOG":
            continue
        title = re.sub(r"^(?:[IVX]+\.|\d+\.)\s+", "", raw).replace("`", "")
        if nice_title(title).lower() not in flat.lower() and title.lower() not in flat.lower():
            problems.append(f"encabezado del .md ausente en el .docx: {raw}")
    for f in sorted(set(re.findall(r"\b([a-z_0-9]+\.sql)\b", md))):
        if f not in flat:
            problems.append(f"migración del .md ausente en el .docx: {f}")
    for line in re.findall(r"^- (20\d\d-\d\d-\d\d) — (.*)$", md, re.M)[-1:]:
        if line[1][:40].replace("`", "") not in flat:
            problems.append(f"última línea del CHANGELOG del .md ausente en el .docx: {line[0]}")
    for term in AUTH_TERMS:
        in_md, in_docx = term in md, term in flat
        if in_md != in_docx:
            problems.append(f"auth: «{term}» {'está en el .md pero no en el .docx' if in_md else 'está en el .docx pero no en el .md'}")
    for pr in problems:
        print("DESINCRONIZADO:", pr, file=sys.stderr)
    print("Coherencia .md <-> .docx: " + ("OK" if not problems else f"{len(problems)} problema(s)"))
    return len(problems)


def main():
    global PENDING_COMMIT, OUT_PATH
    if "--out" in sys.argv:   # p. ej. para validar sin pisar el .docx de la raíz
        OUT_PATH = Path(sys.argv[sys.argv.index("--out") + 1]).resolve()
    if "--pending" in sys.argv:
        i = sys.argv.index("--pending")
        PENDING_COMMIT = sys.argv[i + 1] if i + 1 < len(sys.argv) else "(sin commitear)"
    if not MD_PATH.exists():
        sys.exit(f"No se encontró {MD_PATH}")
    # Pasada 1: recolecta los encabezados para el índice estático del TOC.
    headings = build(Document(), [])
    # Pasada 2: documento real (el campo TOC de Word se refresca al abrir).
    doc = Document()
    build(doc, headings)
    try:
        doc.save(OUT_PATH)
    except PermissionError:
        sys.exit(f"No se pudo escribir {OUT_PATH.name}: ciérralo en Word y vuelve a correr el script.")
    print(f"OK -> {OUT_PATH}  ({len(headings)} encabezados)")
    if verify_sync():
        sys.exit(1)


if __name__ == "__main__":
    main()
